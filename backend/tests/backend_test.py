"""Backend API tests for OBE-LMS STKIP PGRI Pacitan (Go/Gin behind FastAPI proxy)."""
import io
import os
import uuid

import pytest
import requests
from dotenv import dotenv_values

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")
API = f"{BASE_URL}/api"

CREDS = {
    "admin": ("admin@stkippacitan.ac.id", "admin123"),
    "dosen": ("dosen@stkippacitan.ac.id", "dosen123"),
    "mahasiswa": ("ahmad@student.stkippacitan.ac.id", "mahasiswa123"),
    "fajar": ("fajar@student.stkippacitan.ac.id", "mahasiswa123"),
}

TIMEOUT = 60


def login(role):
    email, pw = CREDS[role]
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": pw}, timeout=TIMEOUT)
    if r.status_code != 200:
        pytest.fail(f"Login failed for {role}: {r.status_code} {r.text[:300]}")
    data = r.json()
    assert "token" in data and data["token"]
    return data


@pytest.fixture(scope="session")
def tokens():
    return {r: login(r)["token"] for r in CREDS}


def cli(token):
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json", "Authorization": f"Bearer {token}"})
    return s


@pytest.fixture(scope="session")
def dosen(tokens):
    return cli(tokens["dosen"])


@pytest.fixture(scope="session")
def admin(tokens):
    return cli(tokens["admin"])


@pytest.fixture(scope="session")
def mhs(tokens):
    return cli(tokens["mahasiswa"])


@pytest.fixture(scope="session")
def fajar(tokens):
    return cli(tokens["fajar"])


@pytest.fixture(scope="session")
def course_id(dosen):
    r = dosen.get(f"{API}/courses", timeout=TIMEOUT)
    assert r.status_code == 200, r.text
    courses = r.json()
    assert isinstance(courses, list) and len(courses) >= 1, "dosen has no courses"
    tif = [c for c in courses if c.get("code") == "TIF301"]
    return (tif or courses)[0]["id"]


# ---------------- AUTH ----------------
class TestAuth:
    def test_health(self):
        r = requests.get(f"{API}/health", timeout=TIMEOUT)
        assert r.status_code == 200
        assert r.json()["status"] == "healthy"

    @pytest.mark.parametrize("role,expected", [("admin", "admin"), ("dosen", "dosen"), ("mahasiswa", "mahasiswa")])
    def test_login_roles(self, role, expected):
        data = login(role)
        assert data["user"]["role"] == expected
        assert data["user"]["email"] == CREDS[role][0]
        assert "password_hash" not in data["user"] and "PasswordHash" not in data["user"]

    def test_login_wrong_password(self):
        r = requests.post(f"{API}/auth/login", json={"email": CREDS["admin"][0], "password": "wrong"}, timeout=TIMEOUT)
        assert r.status_code == 401

    def test_login_unknown_email(self):
        r = requests.post(f"{API}/auth/login", json={"email": "nope@x.com", "password": "x"}, timeout=TIMEOUT)
        assert r.status_code == 401

    def test_me(self, dosen):
        r = dosen.get(f"{API}/auth/me", timeout=TIMEOUT)
        assert r.status_code == 200
        assert r.json()["email"] == CREDS["dosen"][0]

    def test_me_no_token(self):
        r = requests.get(f"{API}/auth/me", timeout=TIMEOUT)
        assert r.status_code == 401

    def test_me_bad_token(self):
        r = requests.get(f"{API}/auth/me", headers={"Authorization": "Bearer garbage"}, timeout=TIMEOUT)
        assert r.status_code == 401

    def test_register_new_user(self, admin):
        email = f"TEST_{uuid.uuid4().hex[:8]}@student.test"
        r = requests.post(f"{API}/auth/register", json={
            "name": "TEST_User", "email": email, "password": "secret123", "role": "mahasiswa", "nim": "TEST123"
        }, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        body = r.json()
        assert body["user"]["email"] == email.lower()  # backend normalizes email
        assert body["user"]["role"] == "mahasiswa"
        assert body["token"]
        # verify persistence via /auth/me with new token
        me = requests.get(f"{API}/auth/me", headers={"Authorization": f"Bearer {body['token']}"}, timeout=TIMEOUT)
        assert me.status_code == 200 and me.json()["email"] == email.lower()
        admin.delete(f"{API}/users/{body['user']['id']}", timeout=TIMEOUT)

    def test_register_duplicate_email(self):
        r = requests.post(f"{API}/auth/register", json={
            "name": "Dup", "email": CREDS["dosen"][0], "password": "secret123"
        }, timeout=TIMEOUT)
        assert r.status_code == 400

    def test_register_short_password_rejected(self):
        r = requests.post(f"{API}/auth/register", json={
            "name": "X", "email": f"TEST_{uuid.uuid4().hex[:6]}@x.test", "password": "123"
        }, timeout=TIMEOUT)
        assert r.status_code == 400

    def test_register_self_assign_admin_role(self, admin):
        """SECURITY: public register should not allow role=admin escalation."""
        email = f"TEST_{uuid.uuid4().hex[:8]}@x.test"
        r = requests.post(f"{API}/auth/register", json={
            "name": "TEST_Esc", "email": email, "password": "secret123", "role": "admin"
        }, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        role = r.json()["user"]["role"]
        admin.delete(f"{API}/users/{r.json()['user']['id']}", timeout=TIMEOUT)
        assert role != "admin", "Privilege escalation: public /auth/register granted admin role"


# ---------------- RBAC ----------------
class TestRBAC:
    def test_mahasiswa_blocked_write_endpoints(self, mhs, course_id):
        cases = [
            ("post", f"{API}/courses", {"code": "TEST_X", "name": "X"}),
            ("post", f"{API}/cpmk", {"course_id": course_id, "code": "TEST_C", "description": "d"}),
            ("post", f"{API}/assessments", {"course_id": course_id, "name": "TEST_A", "type": "tugas"}),
            ("post", f"{API}/grades", {"grades": []}),
            ("post", f"{API}/meetings", {"course_id": course_id, "week": 1, "topic": "x"}),
            ("post", f"{API}/remediation", {"course_id": course_id}),
            ("post", f"{API}/users", {"name": "x", "email": "a@b.c", "password": "secret123"}),
        ]
        failures = []
        for method, url, payload in cases:
            r = getattr(mhs, method)(url, json=payload, timeout=TIMEOUT)
            if r.status_code != 403:
                failures.append(f"{url} -> {r.status_code}")
        assert not failures, f"Mahasiswa not blocked (403) on: {failures}"

    def test_dosen_blocked_admin_user_mgmt(self, dosen):
        r = dosen.post(f"{API}/users", json={"name": "x", "email": "TEST_z@x.test", "password": "secret123"}, timeout=TIMEOUT)
        assert r.status_code == 403

    def test_list_users_requires_admin(self, mhs):
        """GET /api/users has no role guard in router - students can enumerate all users."""
        r = mhs.get(f"{API}/users", timeout=TIMEOUT)
        assert r.status_code == 403, f"GET /api/users returned {r.status_code} for mahasiswa (should be 403)"


# ---------------- COURSES ----------------
class TestCourses:
    def test_dosen_sees_own_courses(self, dosen, tokens):
        me = dosen.get(f"{API}/auth/me", timeout=TIMEOUT).json()
        r = dosen.get(f"{API}/courses", timeout=TIMEOUT)
        assert r.status_code == 200
        for c in r.json():
            assert c.get("lecturer_id") == me["id"], f"dosen sees foreign course {c.get('code')}"

    def test_mahasiswa_sees_enrolled_only(self, mhs):
        r = mhs.get(f"{API}/courses", timeout=TIMEOUT)
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_admin_sees_all(self, admin, dosen):
        a = admin.get(f"{API}/courses", timeout=TIMEOUT)
        d = dosen.get(f"{API}/courses", timeout=TIMEOUT)
        assert a.status_code == 200
        assert len(a.json()) >= len(d.json())

    def test_get_course_detail(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}", timeout=TIMEOUT)
        assert r.status_code == 200
        assert r.json()["id"] == course_id

    def test_get_course_not_found(self, dosen):
        r = dosen.get(f"{API}/courses/{uuid.uuid4()}", timeout=TIMEOUT)
        assert r.status_code == 404

    def test_course_crud(self, dosen):
        payload = {"code": f"TEST{uuid.uuid4().hex[:4].upper()}", "name": "TEST_Course", "sks": 3, "semester": 3}
        r = dosen.post(f"{API}/courses", json=payload, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        cid = r.json()["id"]
        assert r.json()["name"] == "TEST_Course"
        g = dosen.get(f"{API}/courses/{cid}", timeout=TIMEOUT)
        assert g.status_code == 200 and g.json()["code"] == payload["code"]
        u = dosen.put(f"{API}/courses/{cid}", json={"name": "TEST_Updated", "sks": 4}, timeout=TIMEOUT)
        assert u.status_code == 200
        g2 = dosen.get(f"{API}/courses/{cid}", timeout=TIMEOUT)
        assert g2.json()["name"] == "TEST_Updated"
        d = dosen.delete(f"{API}/courses/{cid}", timeout=TIMEOUT)
        assert d.status_code in (200, 204)
        assert dosen.get(f"{API}/courses/{cid}", timeout=TIMEOUT).status_code == 404


# ---------------- OBE MAPPING ----------------
class TestOBEMapping:
    def test_list_cpmk_structure(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/cpmk", timeout=TIMEOUT)
        assert r.status_code == 200
        items = r.json()
        assert len(items) >= 1
        first = items[0]
        for key in ("id", "code", "description", "threshold"):
            assert key in first
        assert "cpl_links" in first
        assert "sub_cpmks" in first

    def test_cpmk_create_link_subcpmk_delete(self, dosen, admin, course_id):
        code = f"TEST_CPMK{uuid.uuid4().hex[:4]}"
        r = dosen.post(f"{API}/cpmk", json={"course_id": course_id, "code": code, "description": "TEST desc", "threshold": 70}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        cid = r.json()["id"]
        # link to a CPL
        cpls = dosen.get(f"{API}/cpl", timeout=TIMEOUT)
        assert cpls.status_code == 200 and len(cpls.json()) >= 1
        cpl_id = cpls.json()[0]["id"]
        lk = dosen.post(f"{API}/cpmk/{cid}/cpl", json={"cpl_id": cpl_id, "weight": 1}, timeout=TIMEOUT)
        assert lk.status_code in (200, 201), lk.text
        # sub cpmk
        sc = dosen.post(f"{API}/subcpmk", json={"cpmk_id": cid, "code": "TEST_S1", "description": "sub"}, timeout=TIMEOUT)
        assert sc.status_code in (200, 201), sc.text
        # verify persistence
        listed = dosen.get(f"{API}/courses/{course_id}/cpmk", timeout=TIMEOUT).json()
        mine = [c for c in listed if c["id"] == cid]
        assert mine, "created CPMK not returned in list"
        assert len(mine[0]["cpl_links"] or []) == 1
        assert len(mine[0]["sub_cpmks"] or []) == 1
        # cleanup
        assert dosen.delete(f"{API}/cpmk/{cid}", timeout=TIMEOUT).status_code in (200, 204)


# ---------------- ASSESSMENTS ----------------
class TestAssessments:
    def test_list_assessments_with_links(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/assessments", timeout=TIMEOUT)
        assert r.status_code == 200
        items = r.json()
        assert len(items) >= 1
        assert "cpmk_links" in items[0]
        assert any((a.get("cpmk_links") or []) for a in items), "no assessment tagged to CPMK in seed"

    def test_assessment_create_tag_untag_delete(self, dosen, course_id):
        r = dosen.post(f"{API}/assessments", json={
            "course_id": course_id, "name": "TEST_Assessment", "type": "tugas", "weight": 5, "max_score": 100
        }, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        aid = r.json()["id"]
        cpmk = dosen.get(f"{API}/courses/{course_id}/cpmk", timeout=TIMEOUT).json()[0]
        t = dosen.post(f"{API}/assessments/{aid}/tag", json={"cpmk_id": cpmk["id"], "weight": 1}, timeout=TIMEOUT)
        assert t.status_code in (200, 201), t.text
        listed = dosen.get(f"{API}/courses/{course_id}/assessments", timeout=TIMEOUT).json()
        mine = [a for a in listed if a["id"] == aid][0]
        assert len(mine["cpmk_links"]) == 1
        link_id = mine["cpmk_links"][0]["id"]
        ut = dosen.delete(f"{API}/assessment-cpmk/{link_id}", timeout=TIMEOUT)
        assert ut.status_code in (200, 204)
        listed2 = dosen.get(f"{API}/courses/{course_id}/assessments", timeout=TIMEOUT).json()
        mine2 = [a for a in listed2 if a["id"] == aid][0]
        assert not (mine2.get("cpmk_links") or [])
        assert dosen.delete(f"{API}/assessments/{aid}", timeout=TIMEOUT).status_code in (200, 204)


# ---------------- GRADES ----------------
class TestGrades:
    def test_grade_matrix(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/grades", timeout=TIMEOUT)
        assert r.status_code == 200
        data = r.json()
        assert "assessments" in data and "students" in data
        assert len(data["assessments"]) >= 1
        assert len(data["students"]) >= 1
        assert "scores" in data["students"][0] or "grades" in data["students"][0]

    def test_bulk_save_grades_persists(self, dosen, course_id):
        m = dosen.get(f"{API}/courses/{course_id}/grades", timeout=TIMEOUT).json()
        aid = m["assessments"][0]["id"]
        st = m["students"][0]
        sid = st.get("student_id") or st.get("id")
        original = (st.get("scores") or {}).get(aid)
        r = dosen.post(f"{API}/grades", json={"grades": [{"assessment_id": aid, "student_id": sid, "score": 77.5}]}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        after = dosen.get(f"{API}/courses/{course_id}/grades", timeout=TIMEOUT).json()
        row = [s for s in after["students"] if (s.get("student_id") or s.get("id")) == sid][0]
        assert float((row.get("scores") or {})[aid]) == 77.5
        # restore
        if original is not None:
            dosen.post(f"{API}/grades", json={"grades": [{"assessment_id": aid, "student_id": sid, "score": original}]}, timeout=TIMEOUT)


# ---------------- ANALYTICS ----------------
class TestAnalytics:
    def test_attainment(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/attainment", timeout=TIMEOUT)
        assert r.status_code == 200
        d = r.json()
        assert d["cpmk_meta"] and d["students"] and d["class_cpmk"]
        assert isinstance(d["class_final"], (int, float)) and d["class_final"] > 0
        s0 = d["students"][0]
        assert s0["cpmk"] and "score" in s0["cpmk"][0] and "achieved" in s0["cpmk"][0]

    def test_fajar_at_risk(self, dosen, course_id):
        d = dosen.get(f"{API}/courses/{course_id}/attainment", timeout=TIMEOUT).json()
        fj = [s for s in d["students"] if s["nim"] == "2283207031"]
        assert fj, "Fajar Ramadhan (2283207031) not enrolled"
        assert fj[0]["at_risk"] is True, "Fajar should be AT RISK"

    def test_early_warning(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/early-warning", timeout=TIMEOUT)
        assert r.status_code == 200
        d = r.json()
        assert d["total_at_risk"] >= 1
        assert any(w["nim"] == "2283207031" for w in d["warnings"])
        assert d["warnings"][0]["failing_cpmk"]

    def test_dashboard_roles(self, admin, dosen, mhs):
        a = admin.get(f"{API}/dashboard", timeout=TIMEOUT)
        assert a.status_code == 200 and a.json()["role"] == "admin"
        assert a.json()["total_users"] >= 8
        d = dosen.get(f"{API}/dashboard", timeout=TIMEOUT)
        assert d.status_code == 200 and d.json()["total_courses"] >= 1
        assert d.json()["at_risk"] >= 1
        s = mhs.get(f"{API}/dashboard", timeout=TIMEOUT)
        assert s.status_code == 200 and s.json()["role"] == "mahasiswa"
        assert len(s.json()["courses"]) >= 1


# ---------------- PORTFOLIO ----------------
class TestPortfolio:
    def test_student_own_portfolio(self, mhs, course_id):
        r = mhs.get(f"{API}/students/me/portfolio", params={"course_id": course_id}, timeout=TIMEOUT)
        assert r.status_code == 200
        d = r.json()
        assert d["student"] is not None, "portfolio empty for enrolled student"
        assert d["student"]["cpmk"]
        assert isinstance(d["student"]["final_score"], (int, float))
        assert d["cpmk_meta"]

    def test_student_cannot_read_other_portfolio(self, mhs, dosen, course_id):
        att = dosen.get(f"{API}/courses/{course_id}/attainment", timeout=TIMEOUT).json()
        me = mhs.get(f"{API}/auth/me", timeout=TIMEOUT).json()["id"]
        others = [s for s in att["students"] if s["student_id"] != me]
        assert others
        r = mhs.get(f"{API}/students/{others[0]['student_id']}/portfolio", params={"course_id": course_id}, timeout=TIMEOUT)
        assert r.status_code == 403 or r.json().get("student") is None, \
            "Mahasiswa can read another student's portfolio (IDOR)"


# ---------------- EXPORT ----------------
class TestExport:
    def test_export_siakad(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/export/siakad", timeout=TIMEOUT)
        assert r.status_code == 200, r.text[:300]
        assert "spreadsheet" in r.headers.get("Content-Type", ""), r.headers.get("Content-Type")
        assert r.content[:2] == b"PK"
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        expected = ["NIM", "AKTIVITAS", "HASIL_PROYEK", "QUIZ", "TUGAS", "UTS", "UAS"]
        assert headers[:7] == expected, f"got {headers}"
        assert ws.max_row >= 2

    def test_export_obe(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/export/obe", timeout=TIMEOUT)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("Content-Type", "")
        assert r.content[:2] == b"PK"


# ---------------- MEETINGS / MATERIALS ----------------
class TestMeetings:
    def test_16_meetings_seeded(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/meetings", timeout=TIMEOUT)
        assert r.status_code == 200
        items = r.json()
        assert len(items) == 16, f"expected 16 meetings, got {len(items)}"
        weeks = sorted(m["week"] for m in items)
        assert weeks == list(range(1, 17))

    def test_meeting_and_material_crud(self, dosen, course_id):
        r = dosen.post(f"{API}/meetings", json={"course_id": course_id, "week": 99, "topic": "TEST_Topic"}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        mid = r.json()["id"]
        mat = dosen.post(f"{API}/materials", json={"meeting_id": mid, "title": "TEST_Link", "type": "link", "url": "https://example.com"}, timeout=TIMEOUT)
        assert mat.status_code in (200, 201), mat.text
        matid = mat.json()["id"]
        listed = dosen.get(f"{API}/courses/{course_id}/meetings", timeout=TIMEOUT).json()
        mine = [m for m in listed if m["id"] == mid][0]
        assert len(mine["materials"]) == 1 and mine["materials"][0]["title"] == "TEST_Link"
        assert dosen.delete(f"{API}/materials/{matid}", timeout=TIMEOUT).status_code in (200, 204)
        assert dosen.delete(f"{API}/meetings/{mid}", timeout=TIMEOUT).status_code in (200, 204)

    def test_material_file_upload(self, tokens, dosen, course_id):
        meetings = dosen.get(f"{API}/courses/{course_id}/meetings", timeout=TIMEOUT).json()
        mid = meetings[0]["id"]
        files = {"file": ("TEST_material.txt", b"hello obe", "text/plain")}
        r = requests.post(f"{API}/materials/upload", headers={"Authorization": f"Bearer {tokens['dosen']}"},
                          files=files, data={"meeting_id": mid, "title": "TEST_Upload"}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), f"{r.status_code} {r.text[:300]}"
        body = r.json()
        assert body.get("id")
        url = body.get("url") or body.get("file_path") or ""
        assert url, f"upload response missing url: {body}"
        dosen.delete(f"{API}/materials/{body['id']}", timeout=TIMEOUT)


# ---------------- SUBMISSIONS ----------------
class TestSubmissions:
    def test_student_upload_and_dosen_list(self, tokens, dosen, course_id):
        aid = dosen.get(f"{API}/courses/{course_id}/assessments", timeout=TIMEOUT).json()[0]["id"]
        files = {"file": ("TEST_sub.txt", b"my submission", "text/plain")}
        r = requests.post(f"{API}/submissions/upload", headers={"Authorization": f"Bearer {tokens['mahasiswa']}"},
                          files=files, data={"assessment_id": aid}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), f"{r.status_code} {r.text[:300]}"
        subs = dosen.get(f"{API}/assessments/{aid}/submissions", timeout=TIMEOUT)
        assert subs.status_code == 200
        assert len(subs.json()) >= 1


# ---------------- DISCUSSIONS ----------------
class TestDiscussions:
    def test_create_and_list(self, mhs, dosen, course_id):
        r = mhs.post(f"{API}/discussions", json={"course_id": course_id, "content": "TEST_pertanyaan"}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        did = r.json()["id"]
        lst = dosen.get(f"{API}/courses/{course_id}/discussions", timeout=TIMEOUT)
        assert lst.status_code == 200
        assert any(d["id"] == did for d in lst.json())
        assert dosen.delete(f"{API}/discussions/{did}", timeout=TIMEOUT).status_code in (200, 204)


# ---------------- REMEDIATION ----------------
class TestRemediation:
    def test_remediation_flow(self, dosen, mhs, course_id):
        att = dosen.get(f"{API}/courses/{course_id}/attainment", timeout=TIMEOUT).json()
        fj = [s for s in att["students"] if s["nim"] == "2283207031"][0]
        cpmk_id = fj["cpmk"][0]["cpmk_id"]
        r = dosen.post(f"{API}/remediation", json={
            "course_id": course_id, "student_id": fj["student_id"], "cpmk_id": cpmk_id,
            "action": "TEST_remedial tugas", "status": "planned"
        }, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        rid = r.json()["id"]
        u = dosen.put(f"{API}/remediation/{rid}", json={"status": "done"}, timeout=TIMEOUT)
        assert u.status_code == 200
        lst = dosen.get(f"{API}/courses/{course_id}/remediation", timeout=TIMEOUT).json()
        mine = [x for x in lst if x["id"] == rid]
        assert mine and mine[0]["status"] == "done"
        # student sees only own
        smine = mhs.get(f"{API}/courses/{course_id}/remediation", timeout=TIMEOUT)
        assert smine.status_code == 200
        me = mhs.get(f"{API}/auth/me", timeout=TIMEOUT).json()["id"]
        assert all(x["student_id"] == me for x in smine.json()), "student sees other students' remediation"
        assert dosen.delete(f"{API}/remediation/{rid}", timeout=TIMEOUT).status_code in (200, 204)


# ---------------- ADMIN ----------------
class TestAdmin:
    def test_list_users_filters(self, admin):
        r = admin.get(f"{API}/users", timeout=TIMEOUT)
        assert r.status_code == 200 and len(r.json()) >= 8
        byrole = admin.get(f"{API}/users", params={"role": "mahasiswa"}, timeout=TIMEOUT)
        assert byrole.status_code == 200
        assert byrole.json() and all(u["role"] == "mahasiswa" for u in byrole.json())
        search = admin.get(f"{API}/users", params={"search": "Fajar"}, timeout=TIMEOUT)
        assert search.status_code == 200
        assert any("Fajar" in u["name"] for u in search.json()), search.json()

    def test_admin_created_user_with_uppercase_email_can_login(self, admin):
        """BUG: CreateUser stores email verbatim while Login lowercases it."""
        email = f"TEST_{uuid.uuid4().hex[:6]}@X.test"
        r = admin.post(f"{API}/users", json={"name": "TEST_Case", "email": email, "password": "secret123", "role": "mahasiswa"}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        uid = r.json()["id"]
        lg = requests.post(f"{API}/auth/login", json={"email": email, "password": "secret123"}, timeout=TIMEOUT)
        admin.delete(f"{API}/users/{uid}", timeout=TIMEOUT)
        assert lg.status_code == 200, "admin-created user with mixed-case email cannot login (email not normalized in CreateUser)"

    def test_user_crud(self, admin):
        email = f"test_{uuid.uuid4().hex[:6]}@x.test"
        r = admin.post(f"{API}/users", json={"name": "TEST_Admin_Created", "email": email, "password": "secret123", "role": "dosen", "nidn": "999"}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        uid = r.json()["id"]
        assert r.json()["role"] == "dosen"
        lst = admin.get(f"{API}/users", params={"search": "TEST_Admin_Created"}, timeout=TIMEOUT).json()
        assert any(u["id"] == uid for u in lst)
        # new user can login
        lg = requests.post(f"{API}/auth/login", json={"email": email, "password": "secret123"}, timeout=TIMEOUT)
        assert lg.status_code == 200
        assert admin.delete(f"{API}/users/{uid}", timeout=TIMEOUT).status_code in (200, 204)
        lst2 = admin.get(f"{API}/users", params={"search": "TEST_Admin_Created"}, timeout=TIMEOUT).json()
        assert not any(u["id"] == uid for u in lst2)

    def test_cpl_crud(self, admin):
        r = admin.post(f"{API}/cpl", json={"code": f"TEST_CPL{uuid.uuid4().hex[:4]}", "description": "TEST cpl", "domain": "pengetahuan"}, timeout=TIMEOUT)
        assert r.status_code in (200, 201), r.text
        cid = r.json()["id"]
        lst = admin.get(f"{API}/cpl", timeout=TIMEOUT)
        assert lst.status_code == 200 and any(c["id"] == cid for c in lst.json())
        assert admin.delete(f"{API}/cpl/{cid}", timeout=TIMEOUT).status_code in (200, 204)
        assert not any(c["id"] == cid for c in admin.get(f"{API}/cpl", timeout=TIMEOUT).json())


# ---------------- RPS ----------------
class TestRPS:
    def test_get_and_save_rps(self, dosen, course_id):
        r = dosen.get(f"{API}/courses/{course_id}/rps", timeout=TIMEOUT)
        assert r.status_code == 200
        original = r.json().get("data", "")
        u = dosen.put(f"{API}/courses/{course_id}/rps", json={"data": "{\"TEST\":\"rps\"}"}, timeout=TIMEOUT)
        assert u.status_code == 200, u.text
        g = dosen.get(f"{API}/courses/{course_id}/rps", timeout=TIMEOUT).json()
        assert g.get("data") == "{\"TEST\":\"rps\"}", g
        dosen.put(f"{API}/courses/{course_id}/rps", json={"data": original}, timeout=TIMEOUT)
